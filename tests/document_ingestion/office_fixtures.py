"""Tiny in-memory Office documents for extractor tests.

The docx is assembled by hand from the three parts mammoth needs so the tests do
not depend on a writer library. The pptx uses python-pptx, which markitdown's
``pptx`` extra already installs.
"""

from __future__ import annotations

import io
import zipfile

from PIL import Image
from pptx import Presentation
from pptx.util import Inches

_CONTENT_TYPES = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>
"""

_ROOT_RELS = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>
"""

_DOCUMENT_XML = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p><w:pPr><w:pStyle w:val="Heading1"/></w:pPr><w:r><w:t>Goals</w:t></w:r></w:p>
    <w:p><w:r><w:t>Ship document ingestion for Office formats.</w:t></w:r></w:p>
  </w:body>
</w:document>
"""


def minimal_docx() -> bytes:
    """A valid .docx with one Heading 1 and one paragraph."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", _CONTENT_TYPES)
        archive.writestr("_rels/.rels", _ROOT_RELS)
        archive.writestr("word/document.xml", _DOCUMENT_XML)
    return buffer.getvalue()


def minimal_pptx() -> bytes:
    """A two-slide deck with a speaker note and a table."""
    deck = Presentation()
    first = deck.slides.add_slide(deck.slide_layouts[1])
    first.shapes.title.text = "Ingestion roadmap"
    first.placeholders[1].text = "PDF today"
    first.notes_slide.notes_text_frame.text = "Speaker note: mention the sidecar pattern"
    second = deck.slides.add_slide(deck.slide_layouts[5])
    second.shapes.title.text = "Formats"
    table = second.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(6), Inches(1)).table
    table.cell(0, 0).text = "Format"
    table.cell(0, 1).text = "Engine"
    table.cell(1, 0).text = "pdf"
    table.cell(1, 1).text = "pdf-inspector"
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()


def pptx_with_picture() -> bytes:
    """A one-slide deck whose only shapes are a title and an embedded PNG."""
    png = io.BytesIO()
    Image.new("RGB", (4, 4), color=(200, 30, 30)).save(png, format="PNG")
    png.seek(0)
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    slide.shapes.title.text = "Architecture"
    slide.shapes.add_picture(png, Inches(1), Inches(2), width=Inches(2))
    buffer = io.BytesIO()
    deck.save(buffer)
    return buffer.getvalue()
